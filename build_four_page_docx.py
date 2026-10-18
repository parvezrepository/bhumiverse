import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_p3_four_page_proposal(output_path):
    doc = docx.Document()

    # Page Margins: 0.65 in to ensure it fits comfortably in exactly 4 pages
    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)

    BLACK = RGBColor(0, 0, 0)
    FONT_NAME = "Times New Roman"

    def add_title(text, size=15, bold=True, space_after=2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = BLACK
        return p

    def add_meta(text, size=10, italic=False, bold=False, space_after=2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(size)
        r.font.italic = italic
        r.font.bold = bold
        r.font.color.rgb = BLACK
        return p

    def add_h(text, space_before=6, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = BLACK
        return p

    def add_p(text, bold_prefix=None, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.06
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = FONT_NAME
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = BLACK
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.color.rgb = BLACK
        return p

    def add_bullet(text, bold_prefix=None, space_after=2):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = FONT_NAME
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = BLACK
        r = p.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(12)
        r.font.color.rgb = BLACK
        return p

    def format_table_header(row, titles, widths):
        for i, title in enumerate(titles):
            cell = row.cells[i]
            cell.width = widths[i]
            cell.text = title
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.runs[0]
            run.font.name = FONT_NAME
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = BLACK

    def add_table_row(table, data, widths):
        row = table.add_row()
        for i, text in enumerate(data):
            cell = row.cells[i]
            cell.width = widths[i]
            cell.text = text
            set_cell_margins(cell, top=40, bottom=40, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.runs[0]
            run.font.name = FONT_NAME
            run.font.size = Pt(10)
            run.font.color.rgb = BLACK

    # ================= HEADER BLOCK =================
    add_title("BHUMIverse Project Proposal", size=16, space_after=2)
    add_title("Intelligent Land Record Verification & Fraud Prevention Platform", size=13, bold=False, space_after=3)
    add_meta("LandTech Innovation Hackathon 2026 | Four Page Concept Note", bold=True, size=10.5, space_after=2)
    add_meta("Team name: BHUMIverse | Primary category: P3 Catch Fraud & Double-Selling", bold=True, size=10, space_after=2)
    add_meta("Track: Digital Land Records, Verification and Digitization | Date: 18 October 2026", size=10, space_after=2)
    add_meta("Submitted to: Department of CSE and CSE-TECH, Daffodil International University", italic=True, size=9.5, space_after=6)

    # ================= 1. PROJECT SUMMARY =================
    add_h("1 Project Summary")
    add_p("BHUMIverse is an intelligent land fraud detection and double-selling prevention platform engineered specifically for Bangladesh. Our platform solves Hackathon category P3 ('Catch Fraud & Double-Selling') by identifying fake and tampered documents, catching cases where the same land is sold to more than one buyer, and revealing dangerous record mismatches between deed registration (Sub-Registry) and mutation (AC Land offices). Citizens and buyers can search by parcel or upload deeds to inspect altered text, verify seller decimal limits, and uncover organized fraud syndicates. The platform delivers an explainable Land Verification Report with an instant Risk Score (Low, Medium, High) and clear next steps in both English and Bengali.")

    # ================= 2. SELECTED PROBLEM AND OBJECTIVES =================
    add_h("2 Selected Problem and Objectives")
    add_p("Primary Category: P3 Catch Fraud & Double-Selling (Spot fake documents, double-selling, and registration-mutation mismatches).", bold_prefix="")
    add_p("According to empirical studies by the Law Commission of Bangladesh and the World Bank, over 70% of all civil disputes arise from fraudulent land transactions and ownership conflicts [7]. Citizens face four severe fraud threats: (1) Fake and tampered deeds where photocopy text, dates, or plot sizes are digitally altered; (2) Double-selling where dishonest sellers sell the same plot to multiple buyers before any mutation is recorded; (3) Registration-mutation mismatch where registered deeds never get mutated at the AC Land office, leaving the seller's name in official records to enable second illegal sales; and (4) Fraud syndicates that exploit loopholes across disjointed government offices.")
    add_p("Core Objectives: (a) Spot fake and altered documents using automated Smart Document Diffing; (b) Mathematically detect double-selling and over-selling by tracking seller remaining decimals; (c) Reconcile deed registration records with mutation khatians to flag broken ownership chains; and (d) Map repeat fraud actors and disputed plots on an interactive fraud network graph.")

    # ================= 3. EXISTING SERVICES AND MAIN CONTRIBUTION =================
    add_h("3 Existing Services and Main Contribution")
    add_p("Current government services like e-Porcha, e-Mutation, and online Land Tax operate in silos. Sub-Registry offices record deeds without real-time verification of the seller's live remaining balance at the AC Land office. This gap allows dishonest sellers to sell the same plot multiple times. BHUMIverse connects these separated records into an explained fraud-detection matrix.")

    # Table 1
    t1 = doc.add_table(rows=1, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    w1 = [Inches(1.8), Inches(2.2), Inches(3.2)]
    set_table_borders(t1)
    format_table_header(t1.rows[0], ["Existing Government Baseline", "Limitations in Catching Fraud", "BHUMIverse Platform Contribution"], w1)
    add_table_row(t1, ["e-Porcha & Record Lookup (land.gov.bd)", "Shows individual survey records but cannot verify whether the seller has already sold their share to someone else.", "Builds an automated Transaction Timeline connecting CS-SA-RS-BS and tracking live remaining seller decimals."], w1)
    add_table_row(t1, ["Sub-Registry Deed Registration", "Registers deeds based on presented papers without verifying live mutation records or checking duplicate sales.", "Deterministic Rule Engine cross-checks past deeds, catches duplicate sales, and identifies registered deeds lacking mutation."], w1)
    add_table_row(t1, ["Photocopy Document Verification", "Officers and buyers have no automated tool to detect altered numbers, dates, or names on photocopy deeds.", "Smart Document Diff tool provides side-by-side visual and OCR comparison to highlight tampered text and numbers in red."], w1)
    add_table_row(t1, ["Fraud Syndicate Monitoring", "No government system connects repeat fraudsters across different sub-registries and AC land jurisdictions.", "Fraud Network Graph visually models buyers, sellers, brokers, and disputed plots on an interactive node-link graph."], w1)

    # ================= 4. SYSTEM WORKFLOW =================
    add_h("4 System Workflow")
    add_p("The BHUMIverse platform follows a structured 7-step fraud detection workflow:")
    add_bullet("Citizen accesses the web portal and logs in (or tests pre-loaded benchmark fraud cases).", "Step 1 (Access Platform): ")
    add_bullet("User enters District, Upazila, Mouza, JL Number, Dag, and Khatian number to retrieve historical records.", "Step 2 (Select Land Parcel): ")
    add_bullet("User uploads scanned deeds or porchas. The system generates a cryptographic SHA-256 fingerprint to immediately flag duplicate file submissions.", "Step 3 (Upload Evidence): ")
    add_bullet("User enters proposed purchase area. The engine runs 5 deterministic fraud rules: remaining seller share, survey continuity, registered mortgage, court litigation, and khas/acquisition status.", "Step 4 (Run Fraud Engine): ")
    add_bullet("Results are displayed on an Explainable Radar Chart (covering Authenticity, Continuity, Custody, Encumbrance, and Consistency) and an interactive Fraud Network Graph.", "Step 5 (Visual Diagnostics): ")
    add_bullet("Platform generates a comprehensive Land Verification Report with an overall Risk Score (Low, Medium, High), audit verification hash, and downloadable PDF report.", "Step 6 (Download Audit Report): ")
    add_bullet("Citizens or reviewers can flag contested records for secondary human review and log immutable audit events.", "Step 7 (Case Review & Next Steps): ")
    add_p("A seller owns 10 decimal in Savar Mouza (Dag 10482). In 2021, the seller sold 6 decimal to Buyer A via Deed No. 4821. Buyer A delayed filing for mutation. In 2023, the seller produces an old khatian copy showing 10 decimal and attempts to sell 8 decimal to Buyer B. The BHUMIverse rule engine analyzes past deed transfers and discovers that only 4 decimal remain. When 8 decimal is entered, the system triggers a Critical Double-Selling Alert (6 + 8 = 14 > 10 decimal), stopping the fraudulent sale immediately.", bold_prefix="Real-World Double-Selling Example: ")

    # ================= 5. ARCHITECTURE AND IMPLEMENTATION METHOD =================
    add_h("5 Architecture and Implementation Method")
    add_p("BHUMIverse is engineered with modern, robust, scalable technologies:")
    add_bullet("Next.js 16 (App Router), React 19, TypeScript, and Tailwind CSS. Features Vis-network for interactive fraud syndicate graphs, Chart.js/SVG for explainable radar charts, and full dual-language support (English and Bengali).", "Frontend Web Application: ")
    add_bullet("FastAPI (Python 3.11) with Pydantic v2 schemas. Implements deterministic rule engines, SHA-256 file fingerprinting, and audit logging services.", "Backend REST API: ")
    add_bullet("Supabase (PostgreSQL with Row-Level Security) managing user accounts, parcel metadata, case histories, and audit trails.", "Database & Authentication: ")
    add_bullet("Side-by-side visual difference tool that compares original deed templates against altered uploaded documents to highlight modified text or numbers.", "Smart Document Diff Engine: ")
    add_bullet("Vercel production build for frontend and cloud-ready Dockerized backend architecture.", "Deployment: ")

    # ================= 6. ACTUAL DEVELOPMENT PROGRESS =================
    add_h("6 Actual Development Progress")
    add_p("Status as of 18 October 2026: The BHUMIverse prototype is fully built, tested, and operational. All core features can be tested live.")

    # Table 2
    t2 = doc.add_table(rows=1, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    w2 = [Inches(1.8), Inches(1.8), Inches(3.6)]
    set_table_borders(t2)
    format_table_header(t2.rows[0], ["P3 Fraud Module", "Development Status", "Demonstrated Capabilities in Prototype"], w2)
    add_table_row(t2, ["Double-Selling Detection Engine", "100% Implemented", "Computes remaining seller share balance; alerts when proposed sale exceeds owned decimals."], w2)
    add_table_row(t2, ["Smart Document Diff Tool", "100% Implemented", "Side-by-side deed diffing; visual red-highlighting of altered text and dates; confidence scoring."], w2)
    add_table_row(t2, ["Fraud Network Graph", "100% Implemented", "Interactive physics-based node graph showing linked buyers, sellers, disputed plots, and fraud syndicates."], w2)
    add_table_row(t2, ["Registration-Mutation Mismatch", "100% Implemented", "Cross-checks registered deeds against AC Land mutation khatians to flag non-mutated ghost titles."], w2)
    add_table_row(t2, ["Explainable Radar Chart", "100% Implemented", "5-axis radar chart showing exact scores across Authenticity, Continuity, Custody, Encumbrance, and Consistency."], w2)
    add_table_row(t2, ["Transaction Timeline", "100% Implemented", "Chronological card timeline tracking ownership transfers across CS, SA, RS, and BS survey records."], w2)
    add_table_row(t2, ["Bilingual Support (EN / BN)", "100% Implemented", "Instant one-click toggle between English and Bengali across all pages, forms, and results."], w2)
    add_table_row(t2, ["Supabase Auth & Audit Logs", "100% Implemented", "Secure sign-in, case creation, document uploads, and immutable audit logs."], w2)

    # ================= 7. LEGAL ALIGNMENT AND DATA PROTECTION =================
    add_h("7 Legal Alignment and Data Protection")
    add_p("BHUMIverse is designed in strict compliance with the statutory land laws of Bangladesh:")
    add_bullet("Under Section 53C, any land transfer is legally void if the seller's name is not recorded in the latest khatian (or supported by a valid succession certificate). Under Section 53D, transferring mortgaged property without the lender bank's written consent is illegal. BHUMIverse enforces both checks automatically.", "Transfer of Property Act 1882 (Sections 53C & 53D): ")
    add_bullet("Mandates a complete 25-year chain of title and verified mutation porcha before any deed registration. The BHUMIverse timeline directly verifies this 25-year continuity.", "Registration Act 1908 (Section 52A & 2026 Amendment): ")
    add_bullet("Categorizes deed forgery and cheating as criminal offenses. BHUMIverse identifies tampered photocopies and fraudulent double-selling patterns to protect citizens from criminal deception.", "Penal Code 1860 (Sections 420 & 468 - Fraud & Forgery): ")
    add_bullet("All demonstration data is synthetic and non-confidential. Citizen records are protected with encryption and Row-Level Security (RLS). No unauthorized scraping of government portals is conducted.", "Data Privacy & Ethics: ")

    # ================= 8. LIMITATIONS AND MITIGATION =================
    add_h("8 Limitations and Mitigation")
    # Table 3
    t3 = doc.add_table(rows=1, cols=2)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    w3 = [Inches(3.4), Inches(3.8)]
    set_table_borders(t3)
    format_table_header(t3.rows[0], ["Current Prototype Limitation", "Mitigation Strategy & Future Solution"], w3)
    add_table_row(t3, ["No public API available for live government land databases during hackathon.", "Developed standardized JSON adapters and a realistic 8-parcel synthetic benchmark covering real-world Savar, Dhanmondi, and Mirpur scenarios."], w3)
    add_table_row(t3, ["Handwritten historical deeds in archaic Bangla script are difficult for raw OCR.", "Integrated human-in-the-loop confirmation where the system extracts proposed values and prompts the user to confirm before running rules."], w3)
    add_table_row(t3, ["Automated software cannot replace an official legal court decree or title certificate.", "BHUMIverse explicitly presents itself as an advisory verification tool with clear checklists, advising citizens when to hire a qualified land lawyer."], w3)

    # ================= 9. FUTURE SCOPE IN BANGLADESH =================
    add_h("9 Future Scope in Bangladesh")
    add_bullet("Launch a controlled pilot in one Upazila (e.g., Savar or Gazipur Sadar) in collaboration with the local AC Land office and Sub-Registry.", "1. Upazila-Level Pilot: ")
    add_bullet("Connect directly with the Land Ministry's e-Mutation API, National ID (NID/Porichoy) verification, and the National Land Zoning database.", "2. Official API Integrations: ")
    add_bullet("Deploy offline-capable kiosks at Union Digital Centers (UDCs) and launch an Android mobile app with Bangla voice guidance for rural farmers and overseas expatriates.", "3. Citizen Accessibility: ")

    # ================= 10. EVALUATION AND DELIVERY PLAN =================
    add_h("10 Evaluation and Delivery Plan")
    add_p("For Round 1 of the LandTech Innovation Hackathon 2026, the BHUMIverse team delivers:")
    add_bullet("Complete 4-page concept note and system architecture document focused on P3 fraud prevention.", "1. Project Proposal: ")
    add_bullet("Live, interactive web platform featuring document diffing, fraud network visualization, rule-based verification, and instant report generation.", "2. Working Prototype: ")
    add_bullet("Fully documented open-source repository with setup guides, test suites, and synthetic datasets.", "3. Source Code Repository: ")
    add_bullet("Comprehensive video walkthrough demonstrating all citizen verification flows within 3 minutes.", "4. Video Demonstration: ")

    # ================= 11. TEAM AND MENTOR INFORMATION =================
    add_h("11 Team and Mentor Information")
    add_p("Team Name: BHUMIverse  |  Institution: Daffodil International University (DIU)")

    # Table 4
    t4 = doc.add_table(rows=1, cols=4)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    w4 = [Inches(1.5), Inches(1.9), Inches(2.0), Inches(1.8)]
    set_table_borders(t4)
    format_table_header(t4.rows[0], ["Role", "Name", "Student ID / Affiliation", "Email Contact"], w4)
    add_table_row(t4, ["Student Leader", "MD Parvez Hasan", "ID: 241-15-484 (Dept. of CSE, DIU)", "parvez241-15-484@diu.edu.bd"], w4)
    add_table_row(t4, ["Student Member", "MD FARHAN MAHAMUD", "ID: 241-15-497 (Dept. of CSE, DIU)", "farhan241-15-497@diu.edu.bd"], w4)
    add_table_row(t4, ["Faculty Mentor", "Jawad Ibn Ahad", "Lecturer, Dept. of CSE, DIU", "jawad.cse@diu.edu.bd"], w4)
    add_table_row(t4, ["Industry Mentor", "Fateha Hossain Anushka", "Industry Specialist, Software Eng.", "fateha.industry@mentor.bd"], w4)

    # ================= 12. REFERENCES AND PROTOTYPE =================
    add_h("12 References and Prototype")
    add_bullet("LandTech 2026 Rule Book v1.2, Department of CSE & CSE-TECH, Daffodil International University.", "[1] DIU CSE and CSE-TECH: ")
    add_bullet("e-Mutation, e-Porcha, and Land Development Tax Portals (land.gov.bd).", "[2] Ministry of Land, Bangladesh: ")
    add_bullet("Sections 53C & 53D (Transfer restrictions, khatian recording, mortgage consent).", "[3] Transfer of Property Act 1882: ")
    add_bullet("Section 52A (Mandatory 25-year title chain check and deed format).", "[4] Registration Act 1908 & 2026 Amendment: ")
    add_bullet("Sections 420 & 468 (Cheating and document forgery offenses).", "[5] Penal Code 1860: ")
    add_bullet("https://github.com/parvezrepository/bhumiverse", "[6] BHUMIverse Open-Source Repository: ")
    add_bullet("Studies on Civil Litigation Backlogs, Land Governance, and Title Disputes in Bangladesh.", "[7] Law Commission of Bangladesh & World Bank: ")

    doc.save(output_path)
    print(f"Generated P3 4-page proposal at {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.getcwd(), "BHUMIverse_Project_Proposal.docx")
    build_p3_four_page_proposal(out_file)
